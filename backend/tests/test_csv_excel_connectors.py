"""Track 11.7 — CSV and Excel connectors."""
from __future__ import annotations

import base64
import io

import pytest
from openpyxl import Workbook
from rest_framework import status

from apps.integrations.application.connector_registry import (
    get_connector_registry,
    reset_connector_registry,
)
from apps.integrations.connectors.csv_connector import CsvConnector
from apps.integrations.connectors.excel_connector import ExcelConnector
from apps.integrations.domain.sync import SyncCursor, SyncMode
from tests.conftest import api_data


@pytest.fixture(autouse=True)
def _fresh_registry():
    reset_connector_registry()
    yield
    reset_connector_registry()


def test_registry_lists_csv_and_excel():
    catalog = get_connector_registry().list_types()
    types = {item["connector_type"] for item in catalog}
    assert "csv" in types
    assert "excel" in types
    assert "platform.echo" in types
    assert "rest_api" in types


def test_csv_connector_inline_lifecycle():
    connector = CsvConnector()
    config = {
        "inline_text": "region,revenue\nWest,100\nEast,200\n",
        "has_header": True,
        "table_name": "sales",
    }
    assert connector.test_connection(config=config, credentials={}).ok is True
    schema = connector.discover_schema(config=config, credentials={})
    assert schema.tables[0].name == "sales"
    assert [c.name for c in schema.tables[0].columns] == ["region", "revenue"]
    assert schema.tables[0].columns[1].data_type == "integer"

    batch = connector.extract(
        config=config,
        credentials={},
        cursor=SyncCursor.empty(),
        mode=SyncMode.FULL,
        batch_size=1,
    )
    assert len(batch.rows) == 1
    assert batch.has_more is True
    assert batch.rows[0]["region"] == "West"

    batch2 = connector.extract(
        config=config,
        credentials={},
        cursor=batch.next_cursor,
        mode=SyncMode.FULL,
        batch_size=10,
    )
    assert len(batch2.rows) == 1
    assert batch2.has_more is False


def test_excel_connector_inline_lifecycle():
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["sku", "qty"])
    ws.append(["A1", 5])
    ws.append(["B2", 7])
    buf = io.BytesIO()
    wb.save(buf)
    inline_b64 = base64.b64encode(buf.getvalue()).decode("ascii")

    connector = ExcelConnector()
    config = {"inline_b64": inline_b64, "sheet_name": "Orders", "has_header": True}
    result = connector.test_connection(config=config, credentials={})
    assert result.ok is True
    assert result.details["sheet_name"] == "Orders"

    schema = connector.discover_schema(config=config, credentials={})
    assert schema.tables[0].name == "Orders"
    assert schema.tables[0].row_estimate == 2

    batch = connector.extract(
        config=config,
        credentials={},
        cursor=SyncCursor.empty(),
        mode=SyncMode.FULL,
        batch_size=10,
    )
    assert len(batch.rows) == 2
    assert batch.rows[0]["sku"] == "A1"
    assert batch.rows[1]["qty"] == 7


@pytest.mark.django_db
def test_csv_connection_sync_via_storage_object(auth_client):
    org = api_data(auth_client.post("/api/v1/organizations/", {"name": "CSV Co"}, format="json"))
    ws = api_data(
        auth_client.post(
            "/api/v1/workspaces/",
            {"organization_id": org["id"], "name": "Main"},
            format="json",
        )
    )

    from django.core.files.uploadedfile import SimpleUploadedFile

    csv_bytes = b"city,population\nPune,100\nMumbai,200\n"
    upload = auth_client.post(
        "/api/v1/storage/objects/",
        {
            "organization_id": str(org["id"]),
            "workspace_id": str(ws["id"]),
            "file": SimpleUploadedFile("cities.csv", csv_bytes, content_type="text/csv"),
        },
        format="multipart",
    )
    assert upload.status_code == status.HTTP_201_CREATED
    storage = api_data(upload)

    conn = api_data(
        auth_client.post(
            "/api/v1/connections/",
            {
                "organization_id": org["id"],
                "workspace_id": ws["id"],
                "name": "Cities CSV",
                "connector_type": "csv",
                "config": {
                    "storage_object_id": storage["id"],
                    "has_header": True,
                    "table_name": "cities",
                },
            },
            format="json",
        )
    )

    tested = api_data(auth_client.post(f"/api/v1/connections/{conn['id']}/test/", {}, format="json"))
    assert tested["connection"]["health_status"] == "healthy"

    discovered = api_data(
        auth_client.post(f"/api/v1/connections/{conn['id']}/discover/", {}, format="json")
    )
    assert discovered["schema"]["tables"][0]["name"] == "cities"
    assert discovered["schema"]["tables"][0]["row_estimate"] == 2

    synced = api_data(
        auth_client.post(
            f"/api/v1/connections/{conn['id']}/sync/",
            {"mode": "full", "batch_size": 1},
            format="json",
        )
    )
    assert synced["job"]["status"] == "succeeded"
    assert synced["sync_run"]["rows_loaded"] == 2
    assert synced["connection"]["target_dataset_id"] is not None
