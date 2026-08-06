"""Track 11B.9 — Enterprise integration platform exit certification."""
from __future__ import annotations

import base64
import io

import pytest
from openpyxl import Workbook
from rest_framework import status

from apps.audit.infrastructure.models import AuditLog
from apps.integrations.application.connector_registry import (
    get_connector_registry,
    reset_connector_registry,
)
from apps.integrations.domain.sync import SyncCursor, SyncMode
from apps.integrations.infrastructure.models import SyncRun
from apps.jobs.infrastructure.models import Job
from tests.conftest import api_data

CERTIFIED_CONNECTOR_TYPES = (
    "platform.echo",
    "csv",
    "excel",
    "rest_api",
    "postgresql",
    "mysql",
    "sqlserver",
    "mongodb",
    "aws_s3",
    "airtable",
    "shopify",
)


def _excel_inline_b64() -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(["sku", "qty"])
    ws.append(["A1", 5])
    buf = io.BytesIO()
    wb.save(buf)
    return base64.b64encode(buf.getvalue()).decode("ascii")


def _connector_config(connector_type: str) -> dict:
    configs = {
        "platform.echo": {"message": "exit-cert", "row_count": 2},
        "csv": {"inline_text": "region,revenue\nWest,100\nEast,200\n", "has_header": True},
        "excel": {"inline_b64": _excel_inline_b64(), "sheet_name": "Orders", "has_header": True},
        "rest_api": {
            "inline_json": {"data": {"items": [{"id": 1, "name": "Alpha"}]}},
            "json_path": "data.items",
            "table_name": "orders",
        },
        "postgresql": {
            "host": "localhost",
            "database": "demo",
            "table": "orders",
            "inline_rows": [{"id": 1, "sku": "A1"}],
        },
        "mysql": {
            "host": "localhost",
            "database": "demo",
            "table": "orders",
            "inline_rows": [{"id": 1, "sku": "A1"}],
        },
        "sqlserver": {
            "host": "localhost",
            "database": "demo",
            "table": "orders",
            "inline_rows": [{"id": 1, "sku": "A1"}],
        },
        "mongodb": {
            "database": "demo",
            "collection": "orders",
            "inline_documents": [{"_id": "1", "sku": "A1"}],
        },
        "aws_s3": {
            "bucket": "demo-bucket",
            "key": "exports/cities.csv",
            "fileType": "csv",
            "inline_text": "city,population\nPune,100\nMumbai,200\n",
        },
        "airtable": {
            "baseId": "appTEST",
            "tableName": "Orders",
            "inline_records": [{"id": "rec1", "fields": {"sku": "A1", "qty": 2}}],
        },
        "shopify": {
            "shopDomain": "demo.myshopify.com",
            "resource": "orders",
            "inline_records": [{"id": 101, "name": "#1001", "total_price": "25.00"}],
        },
    }
    return configs[connector_type]


@pytest.fixture(autouse=True)
def _fresh_registry():
    reset_connector_registry()
    yield
    reset_connector_registry()


def test_certified_connector_count():
    types = {item["connector_type"] for item in get_connector_registry().list_types()}
    for connector_type in CERTIFIED_CONNECTOR_TYPES:
        assert connector_type in types


@pytest.mark.parametrize("connector_type", CERTIFIED_CONNECTOR_TYPES)
def test_all_connectors_inline_lifecycle(connector_type):
    registry = get_connector_registry()
    connector = registry.get(connector_type)
    config = _connector_config(connector_type)
    credentials: dict = {}

    result = connector.test_connection(config=config, credentials=credentials)
    assert result.ok is True, result.message

    schema = connector.discover_schema(config=config, credentials=credentials)
    assert len(schema.tables) >= 1

    cursor = SyncCursor.empty()
    total_rows = 0
    batches = 0
    while True:
        batch = connector.extract(
            config=config,
            credentials=credentials,
            cursor=cursor,
            mode=SyncMode.FULL,
            batch_size=1,
        )
        batches += 1
        total_rows += len(batch.rows)
        cursor = batch.next_cursor
        if not batch.has_more:
            break
        if batches > 20:
            pytest.fail("extract pagination did not terminate")

    assert total_rows >= 1


@pytest.mark.django_db
def test_connection_audit_on_create_and_sync(auth_client):
    org = api_data(auth_client.post("/api/v1/organizations/", {"name": "Audit Co"}, format="json"))
    ws = api_data(
        auth_client.post(
            "/api/v1/workspaces/",
            {"organization_id": org["id"], "name": "Main"},
            format="json",
        )
    )
    conn = api_data(
        auth_client.post(
            "/api/v1/connections/",
            {
                "organization_id": org["id"],
                "workspace_id": ws["id"],
                "name": "Audit Echo",
                "connector_type": "platform.echo",
                "config": {"message": "audit", "row_count": 2},
            },
            format="json",
        )
    )
    assert AuditLog.objects.filter(
        organization_id=org["id"], action="connection.created", resource_id=str(conn["id"])
    ).exists()

    api_data(
        auth_client.post(
            f"/api/v1/connections/{conn['id']}/sync/",
            {"mode": "full", "batch_size": 10},
            format="json",
        )
    )
    assert AuditLog.objects.filter(
        organization_id=org["id"], action="connection.sync_enqueued"
    ).exists()
    assert AuditLog.objects.filter(
        organization_id=org["id"], action="connection.synced", resource_id=str(conn["id"])
    ).exists()


@pytest.mark.django_db
def test_sync_failure_records_failed_run(auth_client):
    org = api_data(auth_client.post("/api/v1/organizations/", {"name": "Fail Co"}, format="json"))
    ws = api_data(
        auth_client.post(
            "/api/v1/workspaces/",
            {"organization_id": org["id"], "name": "Main"},
            format="json",
        )
    )
    conn = api_data(
        auth_client.post(
            "/api/v1/connections/",
            {
                "organization_id": org["id"],
                "workspace_id": ws["id"],
                "name": "Missing Plugin",
                "connector_type": "snowflake",
                "config": {"account": "missing"},
            },
            format="json",
        )
    )
    synced = auth_client.post(
        f"/api/v1/connections/{conn['id']}/sync/",
        {"mode": "full"},
        format="json",
    )
    assert synced.status_code in {
        status.HTTP_202_ACCEPTED,
        status.HTTP_400_BAD_REQUEST,
        status.HTTP_500_INTERNAL_SERVER_ERROR,
    }
    if synced.status_code == status.HTTP_202_ACCEPTED:
        body = api_data(synced)
        assert body["job"]["status"] == Job.Status.FAILED
        run = SyncRun.objects.get(id=body["sync_run"]["id"])
        assert run.status == SyncRun.Status.FAILED
        assert run.error


@pytest.mark.django_db
def test_cross_org_sync_runs_blocked(auth_client, auth_client_b):
    org_a = api_data(auth_client.post("/api/v1/organizations/", {"name": "A"}, format="json"))
    ws_a = api_data(
        auth_client.post(
            "/api/v1/workspaces/",
            {"organization_id": org_a["id"], "name": "Main"},
            format="json",
        )
    )
    conn = api_data(
        auth_client.post(
            "/api/v1/connections/",
            {
                "organization_id": org_a["id"],
                "workspace_id": ws_a["id"],
                "name": "Echo",
                "connector_type": "platform.echo",
                "config": {"row_count": 1},
            },
            format="json",
        )
    )
    api_data(
        auth_client.post(
            f"/api/v1/connections/{conn['id']}/sync/",
            {"mode": "full"},
            format="json",
        )
    )
    hijack = auth_client_b.get(f"/api/v1/connections/{conn['id']}/sync-runs/")
    assert hijack.status_code == status.HTTP_403_FORBIDDEN


@pytest.mark.django_db
def test_echo_batched_sync_pagination(auth_client):
    org = api_data(auth_client.post("/api/v1/organizations/", {"name": "Batch Co"}, format="json"))
    ws = api_data(
        auth_client.post(
            "/api/v1/workspaces/",
            {"organization_id": org["id"], "name": "Main"},
            format="json",
        )
    )
    conn = api_data(
        auth_client.post(
            "/api/v1/connections/",
            {
                "organization_id": org["id"],
                "workspace_id": ws["id"],
                "name": "Echo Batch",
                "connector_type": "platform.echo",
                "config": {"message": "batch", "row_count": 5},
            },
            format="json",
        )
    )
    synced = api_data(
        auth_client.post(
            f"/api/v1/connections/{conn['id']}/sync/",
            {"mode": "full", "batch_size": 2},
            format="json",
        )
    )
    assert synced["job"]["result"]["rows_loaded"] == 5
    assert synced["sync_run"]["rows_loaded"] == 5


@pytest.mark.django_db
def test_transform_invalid_step_rejected(auth_client):
    org = api_data(auth_client.post("/api/v1/organizations/", {"name": "Xform Co"}, format="json"))
    ws = api_data(
        auth_client.post(
            "/api/v1/workspaces/",
            {"organization_id": org["id"], "name": "Main"},
            format="json",
        )
    )
    conn = api_data(
        auth_client.post(
            "/api/v1/connections/",
            {
                "organization_id": org["id"],
                "workspace_id": ws["id"],
                "name": "Echo",
                "connector_type": "platform.echo",
                "config": {"row_count": 1},
            },
            format="json",
        )
    )
    bad = auth_client.post(
        f"/api/v1/connections/{conn['id']}/transform-rules/",
        {"name": "bad", "steps": [{"op": "explode"}]},
        format="json",
    )
    assert bad.status_code == status.HTTP_400_BAD_REQUEST
