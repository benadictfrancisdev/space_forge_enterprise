"""excel — single-sheet workbook connector via openpyxl."""
from __future__ import annotations

import io
from typing import Any

from apps.core.exceptions import ValidationError
from apps.integrations.connectors._files import as_bool, load_bytes, rows_to_schema, slice_rows
from apps.integrations.domain.connector import (
    BaseConnector,
    ConnectorCapabilities,
    TestResult,
)
from apps.integrations.domain.schema import DiscoveredSchema
from apps.integrations.domain.sync import ExtractBatch, SyncCursor, SyncMode


def _parse_excel_rows(
    raw: bytes,
    *,
    sheet_name: str | None = None,
    has_header: bool = True,
) -> tuple[str, list[str], list[dict[str, Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ValidationError("openpyxl is required for the excel connector") from exc

    try:
        wb = load_workbook(filename=io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise ValidationError(f"Invalid Excel workbook: {exc}") from exc

    try:
        if sheet_name:
            if sheet_name not in wb.sheetnames:
                raise ValidationError(
                    f"Sheet '{sheet_name}' not found. Available: {', '.join(wb.sheetnames)}"
                )
            ws = wb[sheet_name]
            resolved_sheet = sheet_name
        else:
            ws = wb[wb.sheetnames[0]]
            resolved_sheet = wb.sheetnames[0]

        matrix: list[list[Any]] = []
        for row in ws.iter_rows(values_only=True):
            matrix.append(list(row))
    finally:
        wb.close()

    if not matrix:
        return resolved_sheet, [], []

    # Drop fully empty trailing rows
    while matrix and all(cell is None or str(cell).strip() == "" for cell in matrix[-1]):
        matrix.pop()

    if not matrix:
        return resolved_sheet, [], []

    if has_header:
        headers = [
            str(h).strip() if h is not None and str(h).strip() else f"col_{i+1}"
            for i, h in enumerate(matrix[0])
        ]
        data_rows = matrix[1:]
    else:
        width = max(len(r) for r in matrix)
        headers = [f"col_{i+1}" for i in range(width)]
        data_rows = matrix

    parsed: list[dict[str, Any]] = []
    for row in data_rows:
        item: dict[str, Any] = {}
        for i, header in enumerate(headers):
            val = row[i] if i < len(row) else None
            if val is None:
                item[header] = ""
            else:
                item[header] = val
        # Skip completely empty data rows
        if all(v == "" or v is None for v in item.values()):
            continue
        parsed.append(item)
    return resolved_sheet, headers, parsed


class ExcelConnector(BaseConnector):
    @property
    def connector_type(self) -> str:
        return "excel"

    @property
    def version(self) -> str:
        return "1.0.0"

    @property
    def capabilities(self) -> ConnectorCapabilities:
        return ConnectorCapabilities(
            supports_full_sync=True,
            supports_incremental_sync=True,
            supports_schema_discovery=True,
            supports_scheduled_sync=False,
        )

    def get_config_schema(self) -> dict[str, Any]:
        return {
            "type": "object",
            "properties": {
                "storage_object_id": {
                    "type": "string",
                    "format": "uuid",
                    "title": "Storage object ID",
                    "description": "Uploaded .xlsx file in SpaceForge storage",
                },
                "sheet_name": {
                    "type": "string",
                    "title": "Sheet name",
                    "description": "Defaults to the first sheet",
                },
                "has_header": {"type": "boolean", "default": True, "title": "First row is header"},
                "table_name": {
                    "type": "string",
                    "title": "Logical table name",
                    "description": "Defaults to the sheet name",
                },
            },
            "additionalProperties": True,
        }

    def _load_rows(self, config: dict[str, Any], credentials: dict[str, Any]):
        raw = load_bytes(config=config, credentials=credentials)
        sheet_name = config.get("sheet_name") or None
        if sheet_name is not None:
            sheet_name = str(sheet_name).strip() or None
        has_header = as_bool(config.get("has_header"), True)
        return _parse_excel_rows(raw, sheet_name=sheet_name, has_header=has_header)

    def test_connection(
        self,
        *,
        config: dict[str, Any],
        credentials: dict[str, Any],
    ) -> TestResult:
        try:
            sheet, headers, rows = self._load_rows(config, credentials)
        except ValidationError as exc:
            return TestResult(ok=False, message=str(exc))
        except Exception as exc:  # noqa: BLE001
            return TestResult(ok=False, message=f"Excel read failed: {exc}")
        return TestResult(
            ok=True,
            message=f"Excel readable (sheet={sheet}, {len(rows)} rows, {len(headers)} columns)",
            details={
                "sheet_name": sheet,
                "row_count": len(rows),
                "column_count": len(headers),
                "columns": headers,
            },
        )

    def discover_schema(
        self,
        *,
        config: dict[str, Any],
        credentials: dict[str, Any],
    ) -> DiscoveredSchema:
        sheet, headers, rows = self._load_rows(config, credentials)
        table_name = str(config.get("table_name") or sheet or "excel_data")
        return rows_to_schema(
            table_name=table_name,
            headers=headers,
            sample_rows=rows[:50],
            row_estimate=len(rows),
        )

    def extract(
        self,
        *,
        config: dict[str, Any],
        credentials: dict[str, Any],
        cursor: SyncCursor,
        mode: SyncMode,
        batch_size: int = 100,
    ) -> ExtractBatch:
        _ = mode
        sheet, headers, rows = self._load_rows(config, credentials)
        _ = headers
        offset = int(cursor.state.get("offset", 0))
        chunk, next_offset, has_more = slice_rows(rows, offset=offset, batch_size=batch_size)
        table_name = str(config.get("table_name") or sheet or "excel_data")
        return ExtractBatch(
            rows=tuple(chunk),
            next_cursor=SyncCursor(state={"offset": next_offset}),
            has_more=has_more,
            table_name=table_name,
        )
